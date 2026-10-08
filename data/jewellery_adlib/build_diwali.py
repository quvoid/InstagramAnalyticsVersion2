# -*- coding: utf-8 -*-
"""Diwali 2025 creatives workbook: every post (reels, photos, carousels) each brand published
between 30 Sep and 1 Nov 2025, with collaborators, transcripts and on-image text.

Inputs: diwali2025_reels.json (Reels tab + grid), diwali2025_grid.json (grid-only run),
diwali_transcripts_groq*.json, diwali_image_reads.json.
"""
import json, os, sys, glob, re, collections
from datetime import datetime, timezone, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
from core.kolkata_engine import clean_cell

IST = timezone(timedelta(hours=5, minutes=30))
OUT = ROOT + "/deliverables/Diwali2025_Jewellery_Creatives_30Sep-1Nov_v2.xlsx"
BRANDS = [("GRT Jewellers", "grtjewellers"), ("Malabar Gold", "malabargoldanddiamonds"), ("Tanishq", "tanishqjewellery"),
          ("Joyalukkas", "joyalukkas"), ("Khazana", "khazanajewellery"), ("Lalithaa Jewellery", "lalithaajewellery_"),
          ("Bhima", "bhimajewellers"), ("Thangamayil", "thangamayiljewellery"),
          ("Akshaya Thangamaligai (ATM)", "akshayathangamaligai_atm"), ("Regal Jewellers", "regaljewellers_in")]
ALT = {"bhimajewellers": ["bhima_jewellers_"]}


def load(p):
    try:
        return json.load(open(p, encoding="utf-8"))
    except Exception:
        return {}


SRC = {}
for f in ["diwali2025_grid.json", "diwali2025_reels.json"]:
    for b, r in load(ROOT + "/" + f).items():
        SRC.setdefault(b, {"reels": [], "grid": []})
        SRC[b]["reels"] += r.get("reels") or []
        SRC[b]["grid"] += (r.get("grid") or {}).get("posts") or []
TR = {}
for f in sorted(glob.glob(ROOT + "/diwali_transcripts_groq*.json")):
    TR.update(load(f))
IMG = load(ROOT + "/diwali_image_reads.json")
JUNK = re.compile(r"^(\W*\w{1,4}\W*)\1{5,}")


def transcript(code):
    t = TR.get(code)
    if not t:
        return "", "not transcribed"
    x = (t.get("text") or "").strip()
    if not x or x.startswith("no speech"):
        return "", "Music only - no voiceover"
    if JUNK.match(x) or len(set(x.split())) < max(4, len(x.split()) // 6):
        return "", "Music / lyrics only - no usable voiceover"
    return t.get("lang") or "", x


def posts_of(handle):
    seen, out = set(), []
    for h in [handle] + ALT.get(handle, []):
        s = SRC.get(h) or {}
        for p in s.get("reels", []) + s.get("grid", []):
            if p["code"] in seen:
                continue
            seen.add(p["code"]); out.append(p)
    out.sort(key=lambda p: p.get("taken_at") or 0)
    return out


HDR = PatternFill("solid", fgColor="1F3864"); W = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
SEC = PatternFill("solid", fgColor="D9E1F2"); WRAP = Alignment(wrap_text=True, vertical="top")
LINK = Font(color="0563C1", underline="single")


def put(ws, row, vals, link_cols=()):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row, c, clean_cell(v) if isinstance(v, str) else v)
        cell.alignment = WRAP
        if c in link_cols and isinstance(v, str) and v.startswith("http"):
            cell.hyperlink = v; cell.font = LINK


def header(ws, row, cols, widths):
    for c, h in enumerate(cols, 1):
        x = ws.cell(row, c, h); x.font = W; x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical="center")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


wb = openpyxl.Workbook()
ov = wb.active; ov.title = "Overview"
ov_rows, creators_all = [], []
COLS = ["#", "Post Date", "Format", "Posted By", "Brand Own / Creator", "Collaborators (co-author / tagged / sponsor)",
        "Paid Partnership Label", "Views", "Likes", "Comments", "Post Link", "Caption", "Audio Language",
        "Transcript (auto, Whisper)", "On-Image Text (photos / carousels)", "Offer Shown", "Visual Description"]
WID = [5, 11, 10, 22, 13, 34, 11, 11, 10, 10, 34, 70, 11, 80, 60, 34, 50]
for disp, handle in BRANDS:
    P = posts_of(handle)
    roots = {handle.lower()} | {h.lower() for h in ALT.get(handle, [])}
    ws = wb.create_sheet(disp[:31])
    ws.cell(1, 1, f"{disp} (@{handle}) - every creative posted 30 Sep - 1 Nov 2025 (Diwali / Dhanteras window): "
                  f"{len(P)} posts from the Reels tab and Posts grid.").font = B
    ws.cell(2, 1, "Collaborators = co-authors, tagged accounts and the paid-partnership sponsor on the post. "
                  "Transcripts and on-image text are machine-generated and can mis-hear or mis-read words.")
    header(ws, 4, COLS, WID)
    nc = collections.Counter(); collab = {}
    for i, p in enumerate(P, 1):
        own = (p.get("creator") or "").lower() in roots
        others = []
        for h in (p.get("coauthors") or []) + (p.get("tagged") or []) + (p.get("sponsors") or []) + ([] if own else [p.get("creator")]):
            if h and h.lower() not in roots and h not in others:
                others.append(h)
        for h in others:
            c = collab.setdefault(h, {"n": 0, "views": 0, "paid": 0, "first": p["taken_at"], "links": []})
            c["n"] += 1; c["views"] += p.get("views") or 0; c["paid"] += 1 if p.get("is_paid_partnership") else 0
            c["links"].append(p["post_url"])
        nc[p.get("kind")] += 1
        lang, tr = transcript(p["code"]) if p.get("kind") == "Reel" else ("", "")
        im = IMG.get(p["code"]) or {}
        put(ws, 4 + i, [i, datetime.fromtimestamp(p["taken_at"], IST).strftime("%Y-%m-%d"), p.get("kind"),
                        "@" + (p.get("creator") or ""), "Brand" if own else "Creator / other",
                        ", ".join("@" + h for h in others) or "-",
                        "ON" if p.get("is_paid_partnership") else "OFF",
                        p.get("views") if p.get("views") else ("n/a (photo)" if p.get("kind") != "Reel" else "not shown"),
                        "hidden" if p.get("likes_hidden") else p.get("likes"), p.get("comments"), p["post_url"],
                        p.get("caption_full") or p.get("caption") or "", lang, tr,
                        im.get("text", "") if p.get("kind") != "Reel" else "", im.get("offer", ""), im.get("visual", "")],
            link_cols=(11,))
    ws.freeze_panes = "C5"
    if P:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(COLS))}{4 + len(P)}"
    # creators partnered section
    row = 4 + len(P) + 3
    x = ws.cell(row, 1, f"CREATORS / ACCOUNTS PARTNERED DURING DIWALI 2025 ({len(collab)})"); x.font = Font(bold=True, size=12)
    for c in range(1, 8):
        ws.cell(row, c).fill = SEC
    row += 1
    header(ws, row, ["#", "Account", "Posts", "Total Views", "Paid-Label Posts", "First Post", "Post Links"], WID)
    for j, (h, c) in enumerate(sorted(collab.items(), key=lambda kv: -kv[1]["views"]), 1):
        put(ws, row + j, [j, "@" + h, c["n"], c["views"] or "-", c["paid"],
                          datetime.fromtimestamp(c["first"], IST).strftime("%Y-%m-%d"), " | ".join(c["links"])])
        creators_all.append([disp, "@" + h, c["n"], c["views"] or "-", c["paid"]])
    reels = [p for p in P if p.get("kind") == "Reel"]
    spoken = sum(1 for p in reels if transcript(p["code"])[1] and not transcript(p["code"])[1].startswith(("Music", "not trans")))
    ov_rows.append([disp, "@" + handle, len(P), nc["Reel"], nc["Photo"], nc["Carousel"],
                    sum(1 for p in P if p.get("is_paid_partnership")), len(collab), spoken,
                    sum(1 for p in reels if p["code"] in TR), sum(1 for p in P if p.get("kind") != "Reel" and p["code"] in IMG)])

ov.cell(1, 1, "Diwali 2025 - all creatives posted by each jewellery brand, 30 Sep - 1 Nov 2025 (Reels tab + Posts grid).").font = Font(bold=True, size=13)
ov.cell(2, 1, "One tab per brand with every post, its collaborators, transcript (videos) and on-image text (photos). "
              "Bhima's main account posted nothing in this window. Scanned Oct 2026.")
header(ov, 4, ["Brand", "Instagram", "Total Creatives", "Reels", "Photos", "Carousels", "Paid-Label Posts",
               "Creators / Accounts Partnered", "Reels With Voiceover", "Reels Transcribed", "Images Read"],
       [26, 26, 10, 8, 8, 10, 10, 12, 11, 11, 10])
for i, r in enumerate(ov_rows, 5):
    put(ov, i, r)
import diwali_summary as S
sm = wb.create_sheet("Diwali Summary", 1)
sm.sheet_view.showGridLines = False
for c, w in enumerate([30, 62, 44, 50, 50], 1):
    sm.column_dimensions[get_column_letter(c)].width = w
BIG = Font(bold=True, size=15, color="1F3864"); H2 = Font(bold=True, size=12, color="FFFFFF")
GRT_FILL = PatternFill("solid", fgColor="FFF2CC"); GAP_FILL = PatternFill("solid", fgColor="E2EFDA")


def band(row, text):
    for c in range(1, 6):
        sm.cell(row, c).fill = HDR
    sm.cell(row, 1, text).font = H2
    return row + 1


sm.cell(1, 1, "DIWALI 2025 - WHAT COMPETITORS DID AND WHAT GRT DID (30 Sep - 1 Nov 2025, Instagram)").font = BIG
row = 3
row = band(row, "THE SHORT VERSION")
for t in S.HEADLINE:
    sm.cell(row, 1, "-"); c = sm.cell(row, 2, t); c.alignment = WRAP
    sm.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5); sm.row_dimensions[row].height = 32; row += 1
row += 1
row = band(row, "EACH BRAND AT A GLANCE")
for c, h in enumerate(["Brand", "Main Diwali idea", "Creatives posted", "Biggest reel (views)", "Famous faces used"], 1):
    x = sm.cell(row, c, h); x.font = B; x.fill = SEC
row += 1
for b, n, idea, top, faces in S.GLANCE:
    put(sm, row, [b, idea, n, top, faces])
    if b.startswith("GRT"):
        for c in range(1, 6):
            sm.cell(row, c).fill = GRT_FILL
    row += 1
row += 1
row = band(row, "THEME BY THEME - WHAT COMPETITORS DID, WHAT GRT DID, WHAT GRT CAN DO")
for c, h in enumerate(["Theme", "What competitors did (simple words)", "See it (links)", "What GRT did", "What GRT can do"], 1):
    x = sm.cell(row, c, h); x.font = B; x.fill = SEC
row += 1
for theme, comp, links, grt, todo in S.THEMES:
    top = row
    for j, (label, url) in enumerate(links):
        cell = sm.cell(row + j, 3, label); cell.hyperlink = url; cell.font = LINK; cell.alignment = WRAP
    span = max(1, len(links))
    for col, val in [(1, theme), (2, comp), (4, grt), (5, todo)]:
        cell = sm.cell(top, col, val); cell.alignment = WRAP
        if span > 1:
            sm.merge_cells(start_row=top, start_column=col, end_row=top + span - 1, end_column=col)
    sm.cell(top, 1).font = B
    for r_ in range(top, top + span):
        sm.cell(r_, 4).fill = GRT_FILL; sm.cell(r_, 5).fill = GAP_FILL
        sm.row_dimensions[r_].height = max(30, 150 // span)
    row = top + span + 1
row = band(row, "WHAT GRT POSTED THIS DIWALI")
for c, h in enumerate(["Type", "What it was", "Example link"], 1):
    x = sm.cell(row, c, h); x.font = B; x.fill = SEC
row += 1
for t, d, u in S.GRT_DID:
    put(sm, row, [t, d, u], link_cols=(3,)); row += 1
row += 1
row = band(row, "WHAT GRT WAS MISSING")
for g in S.GAPS:
    sm.cell(row, 1, "-"); c = sm.cell(row, 2, g); c.alignment = WRAP
    sm.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5); row += 1
row += 1
sm.cell(row, 1, "Views are Instagram's public play counts as of Oct 2026. GRT's reels do not show a view count publicly. "
                "Every post behind this summary is listed on the brand tabs.").font = Font(italic=True, color="666666")

import adlib_reference
print("ads library reference rows:", adlib_reference.add_tab(wb, put, header, B, SEC, HDR, W, WRAP, LINK, Font))

cs = wb.create_sheet("All Partnered Creators")
header(cs, 1, ["Brand", "Account", "Posts", "Total Views", "Paid-Label Posts"], [26, 30, 8, 12, 10])
for i, r in enumerate(creators_all, 2):
    put(cs, i, r)
wb.save(OUT)
print(OUT)
for r in ov_rows:
    print(r)
