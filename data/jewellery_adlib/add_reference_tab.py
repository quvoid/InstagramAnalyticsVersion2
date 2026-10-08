# -*- coding: utf-8 -*-
"""Add ONLY the 'Ads library reference only' tab to an existing workbook; every other sheet is left untouched."""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, ROOT); sys.path.insert(0, HERE)
from core.kolkata_engine import clean_cell
import adlib_reference

SRC = ROOT + "/deliverables/Diwali2025_Jewellery_Creatives_30Sep-1Nov_v2.xlsx"
OUT = ROOT + "/deliverables/Diwali2025_Jewellery_Creatives_30Sep-1Nov_v3.xlsx"
HDR = PatternFill("solid", fgColor="1F3864"); W = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
SEC = PatternFill("solid", fgColor="D9E1F2"); WRAP = Alignment(wrap_text=True, vertical="top")
LINK = Font(color="0563C1", underline="single")


def put(ws, row, vals, link_cols=()):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row, c, clean_cell(v) if isinstance(v, str) else v)
        cell.alignment = WRAP
        if c in link_cols and isinstance(v, str) and v.startswith("http"):
            cell.hyperlink = v; cell.font = LINK


def header(ws, row, cols, widths):
    for c, h in enumerate(cols, 1):
        x = ws.cell(row, c, h); x.font = W; x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical="center")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


wb = openpyxl.load_workbook(SRC)
if "Ads library reference only" in wb.sheetnames:
    del wb["Ads library reference only"]
n = adlib_reference.add_tab(wb, put, header, B, SEC, HDR, W, WRAP, LINK, Font)
wb.save(OUT)
print(OUT, n, "rows; sheets:", wb.sheetnames[:3], "...", len(wb.sheetnames))
